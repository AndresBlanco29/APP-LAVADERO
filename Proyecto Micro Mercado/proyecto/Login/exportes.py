import os
from datetime import datetime
from xml.sax.saxutils import escape

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from zyntra_ui import (
    APP_NAME,
    BG,
    BG_HEADER,
    BG_CARD,
    BG_ELEVATED,
    ACCENT,
    CYAN,
    PURPLE,
    PURPLE_DARK,
    TEXT,
    TEXT_MUTED,
    TEXT_DIM,
    BORDER,
    SUCCESS,
    DANGER,
    WARNING,
)


C_BG = colors.HexColor(BG)
C_HEADER = colors.HexColor(BG_HEADER)
C_CARD = colors.HexColor(BG_CARD)
C_ELEVATED = colors.HexColor(BG_ELEVATED)
C_ACCENT = colors.HexColor(ACCENT)
C_CYAN = colors.HexColor(CYAN)
C_PURPLE = colors.HexColor(PURPLE)
C_PURPLE_DARK = colors.HexColor(PURPLE_DARK)
C_TEXT = colors.HexColor(TEXT)
C_MUTED = colors.HexColor(TEXT_MUTED)
C_DIM = colors.HexColor(TEXT_DIM)
C_BORDER = colors.HexColor(BORDER)
C_SUCCESS = colors.HexColor(SUCCESS)
C_DANGER = colors.HexColor(DANGER)
C_WARNING = colors.HexColor(WARNING)


def carpeta_reportes():
    path = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "reportes"))
    os.makedirs(path, exist_ok=True)
    return path


def ruta_reporte(nombre):
    base, ext = os.path.splitext(nombre)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return os.path.join(carpeta_reportes(), f"{base}_{stamp}{ext}")


def _celdas(valor):
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y %H:%M")
    return valor


def _negocio():
    try:
        from Consultas import datos_negocio
        data = datos_negocio() or {}
        return data.get("nombre") or APP_NAME, data.get("direccion") or ""
    except Exception:
        return APP_NAME, ""


def _estilos():
    return {
        "titulo": ParagraphStyle(
            "zyntra_titulo",
            fontName="Helvetica-Bold",
            fontSize=16,
            textColor=C_TEXT,
            leading=20,
            spaceAfter=2,
        ),
        "subtitulo": ParagraphStyle(
            "zyntra_sub",
            fontName="Helvetica",
            fontSize=9,
            textColor=C_MUTED,
            leading=12,
            spaceAfter=8,
        ),
        "seccion": ParagraphStyle(
            "zyntra_seccion",
            fontName="Helvetica-Bold",
            fontSize=10,
            textColor=C_CYAN,
            leading=14,
            spaceBefore=8,
            spaceAfter=6,
        ),
        "cuerpo": ParagraphStyle(
            "zyntra_cuerpo",
            fontName="Helvetica",
            fontSize=9,
            textColor=C_TEXT,
            leading=12,
        ),
        "muted": ParagraphStyle(
            "zyntra_muted",
            fontName="Helvetica",
            fontSize=8,
            textColor=C_DIM,
            leading=11,
        ),
        "kpi_label": ParagraphStyle(
            "zyntra_kpi_label",
            fontName="Helvetica",
            fontSize=7,
            textColor=C_DIM,
            leading=9,
            alignment=TA_LEFT,
        ),
        "kpi_valor": ParagraphStyle(
            "zyntra_kpi_valor",
            fontName="Helvetica-Bold",
            fontSize=13,
            textColor=C_TEXT,
            leading=16,
            alignment=TA_LEFT,
        ),
        "celda": ParagraphStyle(
            "zyntra_celda",
            fontName="Helvetica",
            fontSize=8,
            textColor=C_TEXT,
            leading=11,
            alignment=TA_LEFT,
        ),
        "celda_der": ParagraphStyle(
            "zyntra_celda_der",
            fontName="Helvetica",
            fontSize=8,
            textColor=C_TEXT,
            leading=11,
            alignment=TA_RIGHT,
        ),
        "encabezado": ParagraphStyle(
            "zyntra_th",
            fontName="Helvetica-Bold",
            fontSize=8,
            textColor=C_TEXT,
            leading=11,
            alignment=TA_CENTER,
        ),
        "banner": ParagraphStyle(
            "zyntra_banner",
            fontName="Helvetica-Bold",
            fontSize=14,
            textColor=C_TEXT,
            leading=18,
            alignment=TA_CENTER,
        ),
        "banner_sub": ParagraphStyle(
            "zyntra_banner_sub",
            fontName="Helvetica",
            fontSize=8,
            textColor=C_TEXT,
            leading=11,
            alignment=TA_CENTER,
        ),
        "extra": ParagraphStyle(
            "zyntra_extra",
            fontName="Helvetica",
            fontSize=9,
            textColor=C_MUTED,
            leading=12,
        ),
    }


def _p(texto, estilo, alinear=None):
    valor = escape(str("" if texto is None else texto))
    if alinear is None:
        return Paragraph(valor.replace("\n", "<br/>"), estilo)
    clon = ParagraphStyle(f"{estilo.name}_a", parent=estilo, alignment=alinear)
    return Paragraph(valor.replace("\n", "<br/>"), clon)


def _logo(c, x, y, size=26):
    c.setFillColor(C_PURPLE_DARK)
    c.roundRect(x, y, size, size, 5, fill=1, stroke=0)
    c.setFillColor(colors.white)
    c.setFont("Helvetica-Bold", int(size * 0.55))
    c.drawCentredString(x + size / 2.0, y + size * 0.28, "Z")


def dibujar_marco_pagina(c, doc, titulo):
    w, h = doc.pagesize
    c.saveState()
    c.setFillColor(C_BG)
    c.rect(0, 0, w, h, fill=1, stroke=0)

    c.setFillColor(colors.HexColor("#15102A"))
    c.circle(-40, h + 20, 160, fill=1, stroke=0)
    c.setFillColor(colors.HexColor("#0D1730"))
    c.circle(w + 50, -30, 140, fill=1, stroke=0)

    c.setFillColor(C_HEADER)
    c.rect(0, h - 58, w, 58, fill=1, stroke=0)
    c.setFillColor(C_ACCENT)
    c.rect(0, h - 61, w, 3, fill=1, stroke=0)

    _logo(c, 22, h - 44, 26)
    c.setFillColor(C_PURPLE)
    c.setFont("Helvetica-Bold", 13)
    c.drawString(56, h - 28, "Zyntra")
    zyntra_w = c.stringWidth("Zyntra ", "Helvetica-Bold", 13)
    c.setFillColor(C_CYAN)
    c.drawString(56 + zyntra_w, h - 28, "POS")
    c.setFillColor(C_MUTED)
    c.setFont("Helvetica", 8)
    c.drawString(56, h - 42, str(titulo or "Reporte")[:72])

    ahora = datetime.now().strftime("%d/%m/%Y")
    hora = datetime.now().strftime("%H:%M")
    c.setFillColor(C_DIM)
    c.setFont("Helvetica", 8)
    c.drawRightString(w - 24, h - 26, ahora)
    c.drawRightString(w - 24, h - 40, hora)

    c.setFillColor(C_HEADER)
    c.rect(0, 0, w, 32, fill=1, stroke=0)
    c.setFillColor(C_ACCENT)
    c.rect(0, 32, w, 2, fill=1, stroke=0)
    negocio, _dir = _negocio()
    c.setFillColor(C_DIM)
    c.setFont("Helvetica", 7)
    c.drawString(24, 14, f"{negocio}  ·  Documento generado por {APP_NAME}")
    c.drawRightString(w - 24, 14, f"Pág. {doc.page}")
    c.restoreState()


def _tabla_estilo(filas_datos):
    comandos = [
        ("BACKGROUND", (0, 0), (-1, 0), C_ACCENT),
        ("TEXTCOLOR", (0, 0), (-1, -1), C_TEXT),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.35, C_BORDER),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [C_CARD, C_ELEVATED]),
    ]
    if filas_datos <= 0:
        comandos.append(("BACKGROUND", (0, 0), (-1, 0), C_ACCENT))
    return TableStyle(comandos)


def _tabla_branded(encabezados, filas, ancho_util):
    estilos = _estilos()
    n = max(len(encabezados), 1)
    col_w = [ancho_util / n] * n
    data = [[_p(h, estilos["encabezado"]) for h in encabezados]]
    if filas:
        for row in filas:
            celdas = []
            for i, valor in enumerate(row):
                texto = _celdas(valor)
                estilo = estilos["celda_der"] if _parece_dinero(texto) else estilos["celda"]
                celdas.append(_p(texto, estilo))
            if len(celdas) < n:
                celdas.extend([_p("", estilos["celda"])] * (n - len(celdas)))
            data.append(celdas[:n])
    else:
        data.append([_p("Sin registros", estilos["muted"])] + [_p("", estilos["celda"])] * (n - 1))
    tabla = Table(data, repeatRows=1, colWidths=col_w)
    tabla.setStyle(_tabla_estilo(len(filas or [])))
    return tabla


def _parece_dinero(texto):
    s = str(texto).strip()
    if s.startswith("$") or s.startswith("-$"):
        return True
    limpio = s.replace(",", "").replace(".", "").replace("'", "").replace("-", "").replace(" ", "")
    return limpio.isdigit() and (("'" in s) or s.count(".") >= 1)


def _construir(path, titulo, elementos, pagesize):
    pdf = SimpleDocTemplate(
        path,
        pagesize=pagesize,
        leftMargin=22,
        rightMargin=22,
        topMargin=72,
        bottomMargin=44,
        title=f"{titulo} — {APP_NAME}",
        author=APP_NAME,
    )
    pdf.build(elementos, onFirstPage=lambda c, d: dibujar_marco_pagina(c, d, titulo), onLaterPages=lambda c, d: dibujar_marco_pagina(c, d, titulo))
    return path


def exportar_xlsx(encabezados, filas, nombre_archivo, titulo="", extras=None):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as error:
        raise RuntimeError("Falta el paquete openpyxl para exportar a Excel.") from error

    path = ruta_reporte(nombre_archivo if nombre_archivo.endswith(".xlsx") else f"{nombre_archivo}.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Zyntra POS"
    n = max(len(encabezados), 1)
    fill_bg = PatternFill("solid", fgColor="0B0E14")
    fill_header = PatternFill("solid", fgColor="3077FF")
    fill_accent = PatternFill("solid", fgColor="6B4DE6")
    fill_row_a = PatternFill("solid", fgColor="161A24")
    fill_row_b = PatternFill("solid", fgColor="1C2130")
    fill_extra = PatternFill("solid", fgColor="10131A")
    font_logo = Font(bold=True, size=16, color="A084E8", name="Calibri")
    font_title = Font(bold=True, size=12, color="FFFFFF", name="Calibri")
    font_muted = Font(size=9, color="B8BCC8", name="Calibri")
    font_th = Font(bold=True, size=9, color="FFFFFF", name="Calibri")
    font_td = Font(size=9, color="FFFFFF", name="Calibri")
    thin = Border(
        left=Side(style="thin", color="2A3142"),
        right=Side(style="thin", color="2A3142"),
        top=Side(style="thin", color="2A3142"),
        bottom=Side(style="thin", color="2A3142"),
    )

    def pintar_fila(fila, fill, font=None, alto=18):
        ws.row_dimensions[fila].height = alto
        for col in range(1, n + 1):
            cell = ws.cell(fila, col)
            cell.fill = fill
            if font:
                cell.font = font
            cell.alignment = Alignment(vertical="center")

    pintar_fila(1, fill_bg, font_logo, 26)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
    ws.cell(1, 1, APP_NAME)

    pintar_fila(2, fill_bg, font_title, 20)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n)
    ws.cell(2, 1, titulo or "Reporte")

    pintar_fila(3, fill_bg, font_muted, 16)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=n)
    ws.cell(3, 1, datetime.now().strftime("Generado: %d/%m/%Y %H:%M"))

    pintar_fila(4, fill_accent, alto=6)

    fila = 5
    pintar_fila(fila, fill_header, font_th, 22)
    for col, texto in enumerate(encabezados, 1):
        cell = ws.cell(fila, col, texto)
        cell.fill = fill_header
        cell.font = font_th
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin

    for i, row in enumerate(filas or []):
        fila += 1
        fill = fill_row_a if i % 2 == 0 else fill_row_b
        pintar_fila(fila, fill, font_td, 18)
        for col, valor in enumerate(row, 1):
            cell = ws.cell(fila, col, _celdas(valor))
            cell.fill = fill
            cell.font = font_td
            cell.border = thin
            cell.alignment = Alignment(vertical="center")

    if extras:
        fila += 2
        pintar_fila(fila, fill_extra, Font(bold=True, size=10, color="5EC8FF"), 18)
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=n)
        ws.cell(fila, 1, "Resumen")
        for extra in extras:
            fila += 1
            pintar_fila(fila, fill_extra, font_muted, 16)
            ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=n)
            ws.cell(fila, 1, extra)

    for col in range(1, n + 1):
        ancho = 12
        letter_col = get_column_letter(col)
        for cell in ws[letter_col]:
            if cell.value is not None:
                ancho = max(ancho, min(len(str(cell.value)) + 2, 42))
        ws.column_dimensions[letter_col].width = ancho

    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape" if n > 6 else "portrait"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:5"
    ws.oddHeader.left.text = APP_NAME
    ws.oddFooter.right.text = "Pág. &P"
    wb.save(path)
    return path


def exportar_pdf(titulo, encabezados, filas, nombre_archivo, extras=None, apaisado=False):
    path = ruta_reporte(nombre_archivo if nombre_archivo.endswith(".pdf") else f"{nombre_archivo}.pdf")
    page = landscape(letter) if apaisado else letter
    estilos = _estilos()
    ancho_util = page[0] - 44
    elementos = [
        _p(titulo, estilos["titulo"]),
        _p(datetime.now().strftime("%d/%m/%Y  ·  %H:%M"), estilos["subtitulo"]),
        _tabla_branded(encabezados, filas, ancho_util),
    ]
    if extras:
        elementos.append(Spacer(1, 12))
        elementos.append(_p("Resumen", estilos["seccion"]))
        extra_rows = [[_p(str(extra), estilos["extra"])] for extra in extras]
        extra_tabla = Table(extra_rows, colWidths=[ancho_util])
        extra_tabla.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), C_CARD),
                    ("BOX", (0, 0), (-1, -1), 0.4, C_BORDER),
                    ("LEFTPADDING", (0, 0), (-1, -1), 10),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 10),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ]
            )
        )
        elementos.append(extra_tabla)
    return _construir(path, titulo, elementos, page)


from dinero import formato_dinero as _dinero


def _kpi_card(label, valor, estilos, ancho, color_valor=None):
    valor_style = estilos["kpi_valor"]
    if color_valor is not None:
        valor_style = ParagraphStyle(f"kpi_{label}", parent=estilos["kpi_valor"], textColor=color_valor)
    inner = Table(
        [[_p(label.upper(), estilos["kpi_label"])], [_p(valor, valor_style)]],
        colWidths=[max(ancho - 6, 40)],
    )
    inner.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), C_ELEVATED),
                ("BOX", (0, 0), (-1, -1), 0.4, C_BORDER),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (0, 0), 8),
                ("BOTTOMPADDING", (0, 1), (0, 1), 8),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    return inner


def exportar_cierre_pdf(datos):
    path = ruta_reporte("Resumen_Cierre.pdf")
    estilos = _estilos()
    page = letter
    ancho = page[0] - 44
    descuadre = float(datos.get("descuadre", 0) or 0)
    estado_cierre = str(datos.get("estado") or "").strip()
    if estado_cierre:
        estado = estado_cierre
        if estado_cierre == "CON_DESCUADRE":
            color_estado = C_DANGER
            detalle = str(datos.get("observacion") or "Cierre con diferencias documentadas.")
        elif estado_cierre == "CUADRADO_CON_CORRECCIONES":
            color_estado = C_WARNING
            detalle = "Cuadre con reclasificaciones de medios de pago auditadas."
        else:
            color_estado = C_SUCCESS
            detalle = "Conciliación cuadrada sin diferencias."
    elif abs(descuadre) < 0.5:
        estado = "Cuadre exacto"
        color_estado = C_SUCCESS
        detalle = "El conteo físico coincide con el efectivo esperado a retirar (sin base)."
    elif descuadre > 0:
        estado = f"Sobrante  {_dinero(descuadre)}"
        color_estado = C_SUCCESS
        detalle = "Hay más efectivo del esperado a retirar."
    else:
        estado = f"Faltante  {_dinero(abs(descuadre))}"
        color_estado = C_DANGER
        detalle = "Falta efectivo respecto al valor esperado a retirar (la base no se incluye)."

    banner_style = ParagraphStyle("banner_estado", parent=estilos["banner"], textColor=colors.white)
    banner = Table(
        [[_p(estado, banner_style)], [_p(detalle, estilos["banner_sub"])]],
        colWidths=[ancho],
    )
    banner.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), color_estado),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                ("LEFTPADDING", (0, 0), (-1, -1), 10),
                ("RIGHTPADDING", (0, 0), (-1, -1), 10),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ]
        )
    )

    card_w = ancho / 4.0
    kpis = Table(
        [[
            _kpi_card("Base (gabinete)", _dinero(datos.get("monto_inicial", 0)), estilos, card_w),
            _kpi_card("Ventas totales", _dinero(datos.get("total_ventas", 0)), estilos, card_w, C_CYAN),
            _kpi_card("Esperado a retirar", _dinero(datos.get("esperado", 0)), estilos, card_w),
            _kpi_card("Conteo (sin base)", _dinero(datos.get("conteo_fisico", 0)), estilos, card_w, C_WARNING),
        ]],
        colWidths=[card_w] * 4,
    )
    kpis.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3)]))

    meta = [
        ["Cajero", str(datos.get("cajero") or "-"), "Fecha", str(datos.get("fecha") or "-")],
        ["Tickets", str(datos.get("tickets", 0)), "Artículos", str(datos.get("articulos", 0))],
        ["Inicio de turno", str(datos.get("inicio") or "-"), "Usuario", str(datos.get("usuario") or "-")],
        ["Correcciones", str(datos.get("correcciones", 0)), "Estado", str(datos.get("estado") or "-")],
    ]
    meta_tbl = Table(
        [[_p(c, estilos["muted"] if i % 2 == 0 else estilos["cuerpo"]) for i, c in enumerate(row)] for row in meta],
        colWidths=[ancho * 0.18, ancho * 0.32, ancho * 0.18, ancho * 0.32],
    )
    meta_tbl.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), C_CARD),
                ("BOX", (0, 0), (-1, -1), 0.4, C_BORDER),
                ("LINEBELOW", (0, 0), (-1, -2), 0.3, C_BORDER),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )

    metodos = [
        ["Efectivo", _dinero(datos.get("Efectivo", 0))],
        ["Nequi", _dinero(datos.get("Nequi", 0))],
        ["Tarjeta / Datafono", _dinero(datos.get("Tarjeta", 0))],
        ["Transferencia", _dinero(datos.get("Transferencia", 0))],
        ["Crédito", _dinero(datos.get("Crédito", 0))],
    ]
    movimientos = [
        ["Abonos en efectivo", _dinero(datos.get("abonos", 0))],
        ["Entradas de caja", _dinero(datos.get("entradas", 0))],
        ["Salidas de caja", _dinero(datos.get("salidas", 0))],
    ]

    elementos = [
        _p("Cierre diario", estilos["titulo"]),
        _p("Conciliación de medios de pago y resumen del turno", estilos["subtitulo"]),
        meta_tbl,
        Spacer(1, 12),
        kpis,
        Spacer(1, 10),
        banner,
        Spacer(1, 12),
    ]
    conc = datos.get("medios_conciliacion") or []
    if conc:
        filas_c = [
            [
                str(m.get("etiqueta") or m.get("metodo") or ""),
                _dinero(m.get("sistema")),
                _dinero(m.get("declarado")),
                _dinero(m.get("diferencia")),
            ]
            for m in conc
        ]
        elementos.extend(
            [
                _p("Conciliación por medio", estilos["seccion"]),
                _tabla_branded(["Medio", "Sistema", "Declarado", "Diferencia"], filas_c, ancho),
                Spacer(1, 8),
            ]
        )
    elementos.extend(
        [
            _p("Ventas por método de pago", estilos["seccion"]),
            _tabla_branded(["Método", "Total"], metodos, ancho),
            Spacer(1, 8),
            _p("Movimientos de caja", estilos["seccion"]),
            _tabla_branded(["Tipo", "Monto"], movimientos, ancho),
        ]
    )
    return _construir(path, "Cierre de turno", elementos, page)
